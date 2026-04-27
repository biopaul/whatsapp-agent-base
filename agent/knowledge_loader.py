# agent/knowledge_loader.py — Carga y extrae texto de documentos subidos en WP

import os
import time
import hashlib
import logging
from pathlib import Path

import httpx

logger = logging.getLogger("agentkit")

DOCUMENTS_URL = os.getenv("DOCUMENTS_URL", "")
CACHE_DIR = Path("cache/knowledge")
CACHE_TTL = int(os.getenv("KNOWLEDGE_CACHE_TTL", "3600"))  # 1 hora por defecto

# Limite de caracteres por documento extraido.
# Evita inyectar tablas enormes en cada llamada a Claude.
# Un Excel de 10MB puede generar 500K chars; con 12K es mas que suficiente
# para que Claude entienda el contenido sin desperdiciar tokens.
MAX_CHARS_PER_DOC = int(os.getenv("MAX_DOC_CHARS", "12000"))

# Limite total de caracteres a inyectar en el prompt (todos los docs sumados).
MAX_CHARS_TOTAL = int(os.getenv("MAX_KNOWLEDGE_CHARS", "30000"))

# Cache unificado: texto de todos los docs + lista de docs publicos
_cache_text: str | None = None
_cache_public: list[dict] | None = None
_cache_ts: float = 0.0


def _fetch_doc_list() -> list[dict]:
    """Obtiene la lista de documentos desde la REST API de WordPress."""
    if not DOCUMENTS_URL:
        return []
    try:
        with httpx.Client(timeout=10.0) as client:
            resp = client.get(DOCUMENTS_URL)
        if resp.status_code == 200:
            data = resp.json()
            if isinstance(data, list):
                return data
    except Exception as e:
        logger.error(f"Error al obtener lista de documentos: {e}")
    return []


def _doc_cache_path(doc: dict) -> Path:
    """Ruta local de cache para un documento."""
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    doc_id = doc.get("id", 0)
    name = doc.get("name", "doc")
    ext = Path(name).suffix or ".bin"
    filename = f"{doc_id}_{hashlib.md5(name.encode()).hexdigest()[:8]}{ext}"
    return CACHE_DIR / filename


def _download_doc(doc: dict) -> Path | None:
    """Descarga un documento y lo guarda en cache. Retorna la ruta o None."""
    cache_path = _doc_cache_path(doc)
    if cache_path.exists():
        return cache_path
    url = doc.get("url", "")
    if not url:
        return None
    try:
        with httpx.Client(timeout=30.0) as client:
            resp = client.get(url)
        if resp.status_code == 200:
            cache_path.write_bytes(resp.content)
            return cache_path
        logger.warning(f"No se pudo descargar {doc.get('name')}: HTTP {resp.status_code}")
    except Exception as e:
        logger.error(f"Error descargando {doc.get('name')}: {e}")
    return None


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages.append(text.strip())
        return "\n\n".join(pages)
    except Exception as e:
        logger.error(f"Error leyendo PDF {path.name}: {e}")
        return ""


def _extract_docx(path: Path) -> str:
    try:
        import docx
        doc = docx.Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception as e:
        logger.error(f"Error leyendo DOCX {path.name}: {e}")
        return ""


def _extract_xlsx(path: Path) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        rows = []
        for ws in wb.worksheets:
            ws_rows = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                line = " | ".join(cells).strip(" |")
                if line:
                    ws_rows.append(line)
            if ws_rows:
                rows.append(f"[{ws.title}]\n" + "\n".join(ws_rows))
        return "\n\n".join(rows)
    except Exception as e:
        logger.error(f"Error leyendo XLSX {path.name}: {e}")
        return ""


def _extract_text(path: Path, mime_type: str) -> str:
    """Extrae texto de un documento segun su MIME type."""
    if "pdf" in mime_type:
        return _extract_pdf(path)
    elif "wordprocessingml" in mime_type or "msword" in mime_type:
        return _extract_docx(path)
    elif "spreadsheetml" in mime_type or "ms-excel" in mime_type:
        return _extract_xlsx(path)
    return ""


def _refresh_cache() -> None:
    """Descarga la lista de docs y actualiza ambos caches."""
    global _cache_text, _cache_public, _cache_ts

    if not DOCUMENTS_URL:
        logger.warning("DOCUMENTS_URL no configurado — el agente no tendra acceso a documentos subidos")
        _cache_text = ""
        _cache_public = []
        _cache_ts = time.time()
        return

    docs = _fetch_doc_list()
    if not docs:
        logger.info("No se encontraron documentos en la API")
        _cache_text = ""
        _cache_public = []
        _cache_ts = time.time()
        return

    text_parts: list[str] = []
    public_docs: list[dict] = []
    total_chars = 0

    for doc in docs:
        name = doc.get("name", "documento")
        mime = doc.get("type", "")
        is_public = bool(doc.get("is_public", False))

        # Extraer texto de TODOS los docs (privados y publicos) para contexto IA
        path = _download_doc(doc)
        if not path:
            logger.warning(f"No se pudo descargar: {name}")
            if is_public:
                public_docs.append({"name": name, "url": doc.get("url", ""), "id": doc.get("id")})
            continue

        text = _extract_text(path, mime)
        if not text.strip():
            logger.warning(f"Sin texto extraible: {name} (mime={mime!r}) — verificar formato del archivo")
            if is_public:
                public_docs.append({"name": name, "url": doc.get("url", ""), "id": doc.get("id")})
            continue

        # Aplicar limite por documento
        if len(text) > MAX_CHARS_PER_DOC:
            logger.info(f"Doc truncado: {name} ({len(text)} -> {MAX_CHARS_PER_DOC} chars)")
            text = text[:MAX_CHARS_PER_DOC] + f"\n[... documento truncado a {MAX_CHARS_PER_DOC} caracteres]"

        # Aplicar limite total
        if total_chars + len(text) > MAX_CHARS_TOTAL:
            restante = MAX_CHARS_TOTAL - total_chars
            if restante <= 200:
                logger.info(f"Limite total de conocimiento alcanzado, omitiendo: {name}")
                if is_public:
                    public_docs.append({"name": name, "url": doc.get("url", ""), "id": doc.get("id")})
                continue
            text = text[:restante] + f"\n[... truncado por limite total]"
            logger.info(f"Doc parcialmente incluido: {name} ({restante} chars)")

        label = "[PUBLICO] " if is_public else "[PRIVADO] "
        text_parts.append(f"=== {label}{name} ===\n{text.strip()}")
        total_chars += len(text)
        logger.info(f"Conocimiento cargado: {label}{name} ({len(text)} chars, total={total_chars})")

        # Registrar docs publicos para que el agente pueda enviarlos
        if is_public:
            public_docs.append({"name": name, "url": doc.get("url", ""), "id": doc.get("id")})

    _cache_text = "\n\n".join(text_parts)
    _cache_public = public_docs
    _cache_ts = time.time()
    logger.info(f"Cache de conocimiento actualizado: {len(text_parts)} docs, {total_chars} chars totales")


def get_knowledge_text() -> str:
    """
    Retorna texto extraido de TODOS los documentos (privados y publicos).
    Cache en memoria de 1 hora; archivos se cachean en disco.
    """
    global _cache_text, _cache_ts
    if _cache_text is not None and (time.time() - _cache_ts) < CACHE_TTL:
        return _cache_text
    _refresh_cache()
    return _cache_text or ""


def get_public_docs() -> list[dict]:
    """
    Retorna lista de documentos publicos: [{name, url, id}, ...].
    Estos son los archivos que el agente puede enviar a clientes.
    """
    global _cache_public, _cache_ts
    if _cache_public is not None and (time.time() - _cache_ts) < CACHE_TTL:
        return _cache_public
    _refresh_cache()
    return _cache_public or []


def invalidate_knowledge_cache() -> None:
    """Invalida el cache en memoria (los archivos en disco se mantienen)."""
    global _cache_text, _cache_public, _cache_ts
    _cache_text = None
    _cache_public = None
    _cache_ts = 0.0
